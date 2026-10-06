"""The only places docxium does anything python-docx does not already do."""

from __future__ import annotations

from typing import IO, TYPE_CHECKING, Any

import docx
import docx.document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Inches
from lxml import etree

from . import _core

if TYPE_CHECKING:
    from collections.abc import Iterable, Sequence

    RunFormat = dict[str, Any]
    Run = str | tuple[str, bool | None, bool | None] | tuple[str, RunFormat]
    Para = str | Sequence[Run]

_TR = qn("w:tr")
_TBL = qn("w:tbl")


def _splice(doc: docx.document.Document, xml: str) -> None:
    """Parse a ``<w:body>`` fragment and move its children before the final sectPr."""
    body = doc.element.body
    frag = parse_xml(xml)  # python-docx's parser -> real CT_P / CT_Tbl classes
    sect = body.sectPr
    if sect is None:
        body.extend(list(frag))
        return
    for el in list(frag):
        if el.tag == _TBL:
            # Moving a large subtree between documents in one call scales worse
            # than linearly with its size. A <w:tbl> is therefore moved without
            # its rows, and the rows are re-attached one at a time.
            rows = [c for c in el if c.tag == _TR]
            for tr in rows:
                el.remove(tr)
            sect.addprevious(el)
            for tr in rows:
                el.append(tr)
        else:
            sect.addprevious(el)


def _alignment_xml(alignment: Any) -> str | None:
    """``WD_ALIGN_PARAGRAPH`` member or its name (``"center"``, ``"justify"``) -> ``w:jc`` value."""
    if alignment is None:
        return None
    if isinstance(alignment, str):
        try:
            alignment = WD_ALIGN_PARAGRAPH[alignment.upper()]
        except KeyError:
            try:
                alignment = WD_ALIGN_PARAGRAPH.from_xml(alignment)
            except ValueError:
                raise ValueError(f"unknown alignment {alignment!r}") from None
    return WD_ALIGN_PARAGRAPH.to_xml(alignment)


class DocxiumDocument(docx.document.Document):
    """python-docx's ``Document`` plus a few batch methods backed by Rust.

    Everything else is inherited unchanged.
    """

    def add_paragraphs(
        self,
        paragraphs: Iterable[Para],
        style: Any = None,
        alignment: Any = None,
    ) -> None:
        """Append many paragraphs at once.

        Each item is a ``str`` or an iterable of runs. A run is a ``str``,
        ``(text, bold, italic)`` or ``(text, {...})`` where the dict may hold
        ``bold``, ``italic``, ``underline``, ``strike``, ``all_caps``, ``small_caps``,
        ``superscript``, ``subscript``, ``highlight`` (a color name or
        ``WD_COLOR_INDEX`` member), ``size`` (points), ``color`` (``"RRGGBB"``) and
        ``font``. ``style`` (name or style object) and ``alignment``
        (``WD_ALIGN_PARAGRAPH`` member or name) apply to every paragraph.

        Equivalent to calling ``add_paragraph``/``add_run`` in a loop, but crosses
        into Rust once.
        """
        style_id = self.styles.get_style_id(style, WD_STYLE_TYPE.PARAGRAPH)
        xml = _core.paragraphs_xml(paragraphs, style_id, _alignment_xml(alignment))
        _splice(self, xml)

    def add_headings(
        self,
        texts: Iterable[Para],
        level: int = 1,
        alignment: Any = None,
    ) -> None:
        """Append many headings of one ``level`` (0 is the Title style, 1-9 Heading N).

        Same as ``add_heading(text, level)`` in a loop; items may also be run lists.
        """
        if level not in range(10):
            raise ValueError(f"level must be in range 0-9, got {level}")
        style = "Title" if level == 0 else f"Heading {level}"
        self.add_paragraphs(texts, style=style, alignment=alignment)

    def add_table_from_rows(
        self,
        rows: Iterable[Iterable[Any]],
        style: Any = None,
        *,
        header: bool = False,
        col_widths: Sequence[Any] | None = None,
    ):
        """Append a table built from ``rows`` (non-str cells are ``str()``-ed).

        ``header=True`` renders the first row in bold. ``col_widths`` gives one width
        per column (``docx.shared.Length`` such as ``Inches(2)``, or EMU ints); by
        default the text width is split evenly, as ``add_table`` does.

        Returns the new python-docx ``Table``. Equivalent to ``add_table`` followed
        by setting every ``cell.text``.
        """
        rows = [list(r) for r in rows]
        cols = len(rows[0]) if rows else 0
        if col_widths is None:
            sec = self.sections[-1]  # same width python-docx's add_table uses
            width = Emu(
                (sec.page_width or Inches(8.5))
                - (sec.left_margin or Inches(1))
                - (sec.right_margin or Inches(1))
            )
            twips = [Emu(width // cols).twips if cols else 0] * cols
        else:
            twips = [Emu(int(w)).twips for w in col_widths]
            if len(twips) != cols:
                raise ValueError(
                    f"col_widths has {len(twips)} entries but the table has {cols} columns"
                )
        _splice(self, _core.table_xml(rows, twips, header))
        table = self.tables[-1]
        table.style = style
        return table

    def extract_paragraph_texts(self) -> list[str]:
        """Text of every paragraph in document order, *including* those in tables.

        (``Document.paragraphs`` only lists body-level paragraphs.)
        """
        return _core.extract_paragraph_texts(
            etree.tostring(self.element, encoding="utf-8")
        )


def accelerate(doc: docx.document.Document) -> DocxiumDocument:
    """Upgrade an existing python-docx ``Document`` in place and return it."""
    if isinstance(doc, DocxiumDocument):
        return doc
    if type(doc) is not docx.document.Document:
        raise TypeError(f"cannot accelerate {type(doc).__name__}")
    doc.__class__ = DocxiumDocument
    return doc


def Document(docx: str | IO[bytes] | None = None) -> DocxiumDocument:
    """Same as ``docx.Document``; the result also has the batch methods."""
    import docx as _docx

    return accelerate(_docx.Document(docx))
