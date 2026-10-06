import importlib
import io
import pkgutil
import unittest

import docx
import docx.oxml.ns
import docxium
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

TRICKY = [
    "plain",
    "  leading",
    "trailing  ",
    "a & b < c > d",
    "quote\" apos'",
    "中文 テスト ✓",
    "emoji 😀",
]


def body_xml(doc):
    """Serialized body children, excluding the trailing sectPr."""
    sect = docx.oxml.ns.qn("w:sectPr")
    return [etree.tostring(el) for el in doc.element.body if el.tag != sect]


def spec(n):
    return [
        [
            (
                f"P{i} {TRICKY[i % len(TRICKY)]}",
                True if i % 2 == 0 else None,
                True if i % 3 == 0 else None,
            ),
            TRICKY[(i + 1) % len(TRICKY)],
        ]
        for i in range(n)
    ]


def reference_paragraphs(paras, style=None):
    d = docx.Document()
    for runs in paras:
        p = d.add_paragraph(style=style)
        for r in runs:
            if isinstance(r, str):
                p.add_run(r)
            else:
                run = p.add_run(r[0])
                if r[1] is not None:
                    run.bold = r[1]
                if r[2] is not None:
                    run.italic = r[2]
    return d


class AutoCompat(unittest.TestCase):
    """Derived from python-docx's own module and name listings."""

    def test_every_submodule_is_forwarded(self):
        for m in pkgutil.walk_packages(docx.__path__, "docx."):
            real = importlib.import_module(m.name)
            alias = importlib.import_module("docxium" + m.name[len("docx") :])
            self.assertIs(alias, real, m.name)

    def test_every_public_name_is_forwarded(self):
        ours = {"Document", "__version__"}
        for name in dir(docx):
            if name.startswith("_") or name in ours:
                continue
            self.assertIs(getattr(docxium, name), getattr(docx, name), name)

    def test_from_imports_work(self):
        from docxium.enum.text import WD_ALIGN_PARAGRAPH
        from docxium.shared import Pt

        self.assertIs(WD_ALIGN_PARAGRAPH, docx.enum.text.WD_ALIGN_PARAGRAPH)
        self.assertEqual(Pt(12), docx.shared.Pt(12))

    def test_unknown_attribute(self):
        with self.assertRaises(AttributeError):
            _ = docxium.definitely_not_a_thing


class Passthrough(unittest.TestCase):
    def test_is_a_python_docx_document(self):
        d = docxium.Document()
        self.assertIsInstance(d, docx.document.Document)
        d.add_paragraph("x").add_run("y").bold = True
        d.add_table(2, 2).cell(0, 0).text = "c"
        self.assertEqual(d.paragraphs[0].text, "xy")

    def test_accelerate_existing_document(self):
        d = docx.Document()
        d2 = docxium.accelerate(d)
        self.assertIs(d, d2)
        d.add_paragraphs(["a", "b"])
        self.assertEqual([p.text for p in d.paragraphs], ["a", "b"])
        self.assertIs(docxium.accelerate(d), d)  # idempotent

    def test_open_from_stream(self):
        buf = io.BytesIO()
        d = docxium.Document()
        d.add_paragraphs(["hello"])
        d.save(buf)
        buf.seek(0)
        self.assertEqual(docxium.Document(buf).paragraphs[0].text, "hello")


class Paragraphs(unittest.TestCase):
    def test_xml_identical_to_python_docx(self):
        for n in (0, 1, 50, 777):
            fast = docxium.Document()
            fast.add_paragraphs(spec(n))
            self.assertEqual(body_xml(fast), body_xml(reference_paragraphs(spec(n))), n)

    def test_style(self):
        paras = [[f"h{i}"] for i in range(5)]
        fast = docxium.Document()
        fast.add_paragraphs(paras, style="Heading 1")
        ref = reference_paragraphs(paras, style="Heading 1")
        self.assertEqual(body_xml(fast), body_xml(ref))
        self.assertEqual(fast.paragraphs[0].style.name, "Heading 1")

    def test_plain_strings_and_generators(self):
        d = docxium.Document()
        d.add_paragraphs(t for t in ["a", "b"])
        self.assertEqual([p.text for p in d.paragraphs], ["a", "b"])

    def test_roundtrip_and_continued_editing(self):
        d = docxium.Document()
        d.add_paragraphs(spec(30))
        buf = io.BytesIO()
        d.save(buf)
        buf.seek(0)
        d2 = docx.Document(buf)  # opened by plain python-docx
        self.assertEqual(
            [p.text for p in d2.paragraphs],
            [p.text for p in reference_paragraphs(spec(30)).paragraphs],
        )
        d2.add_paragraph("after")
        self.assertEqual(d2.paragraphs[-1].text, "after")

    def test_order_with_existing_content(self):
        d = docxium.Document()
        d.add_paragraph("first")
        d.add_paragraphs(["mid"])
        d.add_paragraph("last")
        self.assertEqual([p.text for p in d.paragraphs], ["first", "mid", "last"])

    def test_bad_input(self):
        with self.assertRaises(TypeError):
            docxium.Document().add_paragraphs([123])


RUN_FORMAT_CASES = [
    {"bold": True},
    {"italic": True, "underline": True},
    {"strike": True, "size": 10.5},
    {"font": "Arial"},
    {"font": 'Fo"nt & <x>'},
    {"color": "ff0000"},
    {"color": "#00AAFF", "size": 11.3},
    {"all_caps": True},
    {"small_caps": True},
    {"superscript": True},
    {"subscript": True},
    {"highlight": "yellow"},
    {"highlight": "darkBlue"},
    {"highlight": "green"},
    {
        "bold": True,
        "italic": True,
        "underline": True,
        "strike": True,
        "all_caps": True,
        "small_caps": True,
        "superscript": True,
        "highlight": "lightGray",
        "size": 24,
        "color": "123ABC",
        "font": "Times New Roman",
    },
    {"bold": None, "size": None, "color": None, "font": None, "highlight": None},
]


def reference_run(text, fmt):
    d = docx.Document()
    r = d.add_paragraph().add_run(text)
    if fmt.get("font"):
        r.font.name = fmt["font"]
    if fmt.get("bold"):
        r.bold = True
    if fmt.get("italic"):
        r.italic = True
    if fmt.get("all_caps"):
        r.font.all_caps = True
    if fmt.get("small_caps"):
        r.font.small_caps = True
    if fmt.get("strike"):
        r.font.strike = True
    if fmt.get("color"):
        r.font.color.rgb = RGBColor.from_string(fmt["color"].lstrip("#").upper())
    if fmt.get("size"):
        r.font.size = Pt(fmt["size"])
    if fmt.get("highlight"):
        r.font.highlight_color = WD_COLOR_INDEX.from_xml(fmt["highlight"])
    if fmt.get("underline"):
        r.font.underline = True
    if fmt.get("superscript"):
        r.font.superscript = True
    if fmt.get("subscript"):
        r.font.subscript = True
    return d


class RunFormatting(unittest.TestCase):
    def test_xml_identical_to_python_docx(self):
        for fmt in RUN_FORMAT_CASES:
            fast = docxium.Document()
            fast.add_paragraphs([[("fmt", fmt)]])
            self.assertEqual(body_xml(fast), body_xml(reference_run("fmt", fmt)), fmt)

    def test_readable_through_python_docx_api(self):
        d = docxium.Document()
        fmt = {
            "bold": True,
            "size": 14,
            "color": "FF0000",
            "font": "Arial",
            "underline": True,
            "superscript": True,
            "all_caps": True,
            "highlight": "yellow",
        }
        d.add_paragraphs([[("x", fmt)]])
        r = d.paragraphs[0].runs[0]
        self.assertTrue(r.bold)
        self.assertEqual(r.font.size.pt, 14)
        self.assertEqual(str(r.font.color.rgb), "FF0000")
        self.assertEqual(r.font.name, "Arial")
        self.assertTrue(r.underline)
        self.assertTrue(r.font.superscript)
        self.assertTrue(r.font.all_caps)
        self.assertEqual(r.font.highlight_color, WD_COLOR_INDEX.YELLOW)

    def test_highlight_spellings(self):
        # An OOXML name, a snake_case name and a WD_COLOR_INDEX member are accepted.
        # OOXML "green" maps to BRIGHT_GREEN and "darkGreen" maps to GREEN.
        for given, expected in [
            ("yellow", WD_COLOR_INDEX.YELLOW),
            ("dark_blue", WD_COLOR_INDEX.DARK_BLUE),
            ("DARKBLUE", WD_COLOR_INDEX.DARK_BLUE),
            ("green", WD_COLOR_INDEX.BRIGHT_GREEN),
            (WD_COLOR_INDEX.GREEN, WD_COLOR_INDEX.GREEN),
            (WD_COLOR_INDEX.TURQUOISE, WD_COLOR_INDEX.TURQUOISE),
        ]:
            d = docxium.Document()
            d.add_paragraphs([[("x", {"highlight": given})]])
            got = d.paragraphs[0].runs[0].font.highlight_color
            self.assertEqual(got, expected, given)

    def test_mixed_run_styles_in_one_paragraph(self):
        d = docxium.Document()
        d.add_paragraphs(
            [
                [
                    "plain ",
                    ("bold", True, None),
                    ("big", {"size": 20}),
                    ("legacy", None, True),
                ]
            ]
        )
        runs = d.paragraphs[0].runs
        self.assertEqual([r.text for r in runs], ["plain ", "bold", "big", "legacy"])
        self.assertEqual([r.bold for r in runs], [None, True, None, None])
        self.assertEqual(runs[2].font.size.pt, 20)
        self.assertTrue(runs[3].italic)

    def test_bad_formats(self):
        d = docxium.Document()
        for bad in (
            {"colour": "red"},
            {"color": "red"},
            {"color": "12345"},
            {"size": 0},
            {"size": -3},
            {"size": 5000},
            {"highlight": "chartreuse"},
            {"highlight": WD_COLOR_INDEX.AUTO},
            {"superscript": True, "subscript": True},
        ):
            with self.assertRaises(ValueError, msg=bad):
                d.add_paragraphs([[("x", bad)]])
        with self.assertRaises(TypeError):
            d.add_paragraphs([[("x", "not a dict")]])


class ParagraphAlignment(unittest.TestCase):
    PARAS = (["left"], ["mid ", ("dle", True, None)], "plain")

    @staticmethod
    def reference(paras, alignment, style=None):
        d = docx.Document()
        for runs in paras:
            p = d.add_paragraph(style=style)
            for r in [runs] if isinstance(runs, str) else runs:
                if isinstance(r, str):
                    p.add_run(r)
                else:
                    run = p.add_run(r[0])
                    if r[1] is not None:
                        run.bold = r[1]
            p.alignment = alignment
        return d

    def test_xml_identical_to_python_docx(self):
        for member in (
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            WD_ALIGN_PARAGRAPH.DISTRIBUTE,
        ):
            fast = docxium.Document()
            fast.add_paragraphs(self.PARAS, alignment=member)
            ref = self.reference(self.PARAS, member)
            self.assertEqual(body_xml(fast), body_xml(ref), member)

    def test_with_style_pstyle_precedes_jc(self):
        fast = docxium.Document()
        fast.add_paragraphs(self.PARAS, style="Heading 2", alignment="center")
        ref = self.reference(self.PARAS, WD_ALIGN_PARAGRAPH.CENTER, style="Heading 2")
        self.assertEqual(body_xml(fast), body_xml(ref))

    def test_names(self):
        for name, member in [
            ("center", WD_ALIGN_PARAGRAPH.CENTER),
            ("RIGHT", WD_ALIGN_PARAGRAPH.RIGHT),
            ("justify", WD_ALIGN_PARAGRAPH.JUSTIFY),
            ("both", WD_ALIGN_PARAGRAPH.JUSTIFY),  # raw OOXML value
        ]:
            d = docxium.Document()
            d.add_paragraphs(["x"], alignment=name)
            self.assertEqual(d.paragraphs[0].alignment, member, name)

    def test_none_adds_nothing(self):
        d = docxium.Document()
        d.add_paragraphs(["x"])
        self.assertIsNone(d.paragraphs[0].alignment)
        self.assertNotIn(b"<w:pPr", etree.tostring(d.paragraphs[0]._p))

    def test_unknown_alignment(self):
        with self.assertRaises(ValueError):
            docxium.Document().add_paragraphs(["x"], alignment="diagonal")


class Headings(unittest.TestCase):
    def test_xml_identical_to_add_heading_for_every_level(self):
        for level in range(10):
            texts = [f"H{level}a", f"H{level} b & c"]
            fast = docxium.Document()
            fast.add_headings(texts, level=level)
            ref = docx.Document()
            for t in texts:
                ref.add_heading(t, level)
            self.assertEqual(body_xml(fast), body_xml(ref), level)

    def test_default_level_and_alignment(self):
        d = docxium.Document()
        d.add_headings(["a", "b"], alignment="center")
        self.assertEqual(
            [p.style.name for p in d.paragraphs], ["Heading 1", "Heading 1"]
        )
        self.assertEqual(d.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_run_lists_allowed(self):
        d = docxium.Document()
        d.add_headings([["Part ", ("I", {"color": "FF0000"})]], level=2)
        self.assertEqual(d.paragraphs[0].text, "Part I")
        self.assertEqual(d.paragraphs[0].style.name, "Heading 2")

    def test_bad_level(self):
        for level in (-1, 10, 99):
            with self.assertRaises(ValueError, msg=level):
                docxium.Document().add_headings(["x"], level=level)


TABLE_ROWS = [
    [f"r{r}c{c} {TRICKY[(r + c) % len(TRICKY)]}" for c in range(5)] for r in range(30)
]


def reference_table(rows, *, header=False, widths=None, style=None):
    d = docx.Document()
    t = d.add_table(len(rows), len(rows[0]))
    t.style = style
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.text = text
            if header and r == 0 and text:
                cell.paragraphs[0].runs[0].bold = True
    if widths is not None:
        for c, w in enumerate(widths):
            t.columns[c].width = w
            for row in t.rows:
                row.cells[c].width = w
    return d


class Tables(unittest.TestCase):
    def test_xml_identical_to_python_docx(self):
        fast = docxium.Document()
        fast.add_table_from_rows(TABLE_ROWS)
        self.assertEqual(body_xml(fast), body_xml(reference_table(TABLE_ROWS)))

    def test_header_bold_identical_to_python_docx(self):
        rows = [["Name", "Item", "Qty"], ["a & b", "x", "1"], ["c", "y", "2"]]
        fast = docxium.Document()
        t = fast.add_table_from_rows(rows, header=True)
        self.assertEqual(body_xml(fast), body_xml(reference_table(rows, header=True)))
        self.assertTrue(t.cell(0, 0).paragraphs[0].runs[0].bold)
        self.assertIsNone(t.cell(1, 0).paragraphs[0].runs[0].bold)

    def test_empty_header_cell(self):
        t = docxium.Document().add_table_from_rows([["", "h"], ["a", "b"]], header=True)
        self.assertEqual([c.text for c in t.rows[0].cells], ["", "h"])
        self.assertTrue(t.cell(0, 1).paragraphs[0].runs[0].bold)

    def test_col_widths_identical_to_python_docx(self):
        widths = [Inches(1), Inches(2.5), Inches(0.75), Inches(1.25), Inches(3)]
        fast = docxium.Document()
        t = fast.add_table_from_rows(TABLE_ROWS, col_widths=widths)
        ref = reference_table(TABLE_ROWS, widths=widths)
        self.assertEqual(body_xml(fast), body_xml(ref))
        self.assertEqual([c.width for c in t.columns], widths)
        self.assertEqual(t.cell(3, 1).width, widths[1])

    def test_col_widths_accept_plain_emu_ints(self):
        t = docxium.Document().add_table_from_rows(
            [["a", "b"]], col_widths=[914400, 1828800]
        )
        self.assertEqual([c.width for c in t.columns], [Inches(1), Inches(2)])

    def test_header_widths_and_style_together(self):
        widths = [Inches(2), Inches(1)]
        rows = [["h1", "h2"], ["a", "b"]]
        fast = docxium.Document()
        t = fast.add_table_from_rows(rows, "Table Grid", header=True, col_widths=widths)
        ref = reference_table(rows, header=True, widths=widths, style="Table Grid")
        self.assertEqual(body_xml(fast), body_xml(ref))
        self.assertEqual(t.style.name, "Table Grid")

    def test_wrong_number_of_col_widths(self):
        for widths in ([Inches(1)], [Inches(1)] * 3):
            with self.assertRaises(ValueError):
                docxium.Document().add_table_from_rows([["a", "b"]], col_widths=widths)

    def test_header_and_widths_options_are_keyword_only(self):
        with self.assertRaises(TypeError):
            docxium.Document().add_table_from_rows([["a"]], None, True)

    def test_returns_table_and_style(self):
        d = docxium.Document()
        t = d.add_table_from_rows(TABLE_ROWS, style="Table Grid")
        self.assertEqual(t.style.name, "Table Grid")
        self.assertEqual(len(t.rows), 30)
        self.assertEqual(len(t.columns), 5)
        self.assertEqual(t.cell(29, 4).text, TABLE_ROWS[29][4])

    def test_non_str_cells_and_empty(self):
        t = docxium.Document().add_table_from_rows([[1, 2.5, None, ""]])
        self.assertEqual([c.text for c in t.rows[0].cells], ["1", "2.5", "None", ""])

    def test_ragged_rows_rejected(self):
        with self.assertRaises(ValueError):
            docxium.Document().add_table_from_rows([["a", "b"], ["c"]])

    def test_empty_inputs_do_not_crash(self):
        d = docxium.Document()
        d.add_paragraphs([])
        d.add_headings([])
        d.add_table_from_rows([])
        d.add_table_from_rows([], header=True, col_widths=[])

    def test_continued_editing(self):
        d = docxium.Document()
        d.add_table_from_rows(TABLE_ROWS)
        d.tables[0].cell(0, 0).text = "edited"
        d.add_paragraph("after")
        self.assertEqual(d.tables[0].cell(0, 0).text, "edited")


class Extract(unittest.TestCase):
    def test_matches_python_docx_and_includes_tables(self):
        d = docxium.Document()
        d.add_paragraphs(spec(50))
        d.add_table_from_rows([["a", "b"], ["c & d", "中"]])
        self.assertEqual(
            d.extract_paragraph_texts(),
            [p.text for p in d.paragraphs] + ["a", "b", "c & d", "中"],
        )

    def test_numeric_char_refs(self):
        from docxium import _core

        d = docxium.Document()
        d.add_paragraphs(["中文 ✓ 😀"])
        ascii_xml = etree.tostring(d.element)  # lxml default: &#20013; style refs
        self.assertEqual(_core.extract_paragraph_texts(ascii_xml), ["中文 ✓ 😀"])

    def test_empty_paragraphs(self):
        d = docxium.Document()
        d.add_paragraph("")
        d.add_paragraph("x")
        self.assertEqual(d.extract_paragraph_texts(), ["", "x"])

    def test_on_documents_not_made_by_docxium(self):
        d = docxium.accelerate(reference_paragraphs(spec(20)))
        self.assertEqual(d.extract_paragraph_texts(), [p.text for p in d.paragraphs])


if __name__ == "__main__":
    unittest.main()
